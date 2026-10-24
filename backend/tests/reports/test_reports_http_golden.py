import csv
import io
import json
import math
import os
import re
from collections.abc import Iterable
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any

import fitz
import pytest
import time_machine
from httpx import AsyncClient, Response
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.domains.identity.models import User
from app.domains.identity.tokens import create_token
from app.domains.market_data.history.models import AssetPriceHistory
from app.domains.market_data.prices.quote import PriceQuote
from app.domains.market_data.prices.service import price_feed

GOLDEN_DIR = Path(__file__).parent / "golden"
UPDATE = os.environ.get("GOLDEN_UPDATE") == "1"
UUID_RE = re.compile(r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}")
SET_RE = re.compile(r"\{('[^']*'(?:, '[^']*')*)\}")
FROZEN_NOW = datetime(2026, 10, 17, 12, 0, tzinfo=UTC)
HISTORY_DAYS = 420
PDF = "application/pdf"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

ASSETS = [
    ("BTC", "crypto", "Binance", "60000"),
    ("ETH", "crypto", "Kraken", "3000"),
    ("SOL", "crypto", "Tangem", "150"),
    ("ARB", "crypto", "Binance", "0.8"),
    ("UNI", "crypto", "Kraken", "9"),
    ("USDC", "crypto", "Binance", "0.92"),
    ("DOGE", "crypto", "Binance", "0.15"),
    ("AAPL", "stock", "Trade Republic", "200"),
]
QUOTES = {symbol: Decimal(price) for symbol, _, _, price in ASSETS}
TRANSACTIONS = [
    ("BTC", "buy", "0.5", "30000", "10", "2025-01-10T10:00:00Z"),
    ("BTC", "buy", "0.25", "50000", "5", "2025-08-01T10:00:00Z"),
    ("BTC", "sell", "0.1", "70000", "7", "2025-11-15T10:00:00Z"),
    ("ETH", "buy", "4", "2000", "4", "2025-02-15T10:00:00Z"),
    ("ETH", "sell", "1", "3500", "3", "2026-03-01T10:00:00Z"),
    ("SOL", "buy", "50", "100", "2", "2025-04-01T10:00:00Z"),
    ("ARB", "buy", "1000", "1.2", "1", "2025-03-01T10:00:00Z"),
    ("UNI", "buy", "200", "7", "1", "2025-06-01T10:00:00Z"),
    ("USDC", "buy", "2000", "0.93", "0", "2025-01-05T10:00:00Z"),
    ("DOGE", "buy", "10000", "0.1", "1", "2025-09-01T10:00:00Z"),
    ("AAPL", "buy", "20", "170", "1", "2025-05-20T10:00:00Z"),
]
SERIES = {
    "BTC": (40000, 0.0012, 0.08, 23, 0.0),
    "ETH": (2200, 0.0009, 0.11, 17, 1.0),
    "SOL": (90, 0.0015, 0.15, 13, 2.0),
    "ARB": (1.1, -0.0008, 0.2, 11, 0.3),
    "UNI": (7, 0.0004, 0.18, 19, 1.7),
    "USDC": (0.92, 0.0, 0.002, 7, 0.0),
    "DOGE": (0.1, 0.001, 0.25, 9, 2.4),
    "AAPL": (165, 0.0004, 0.04, 31, 0.5),
    "PAXG": (2000, 0.0003, 0.02, 41, 1.5),
    "SPY": (480, 0.0003, 0.03, 29, 0.7),
    "^GSPC": (5000, 0.0003, 0.03, 29, 0.6),
}
TARGETS = {"L1": 0.5, "L2": 0.1, "DeFi": 0.1, "Stable": 0.2, "Meme": 0.1}

GET_ROUTES = [
    ("performance_pdf", "/api/v1/reports/performance/pdf"),
    ("performance_excel", "/api/v1/reports/performance/excel"),
    ("tax_summary_2026", "/api/v1/reports/tax/2026/summary"),
    ("tax_summary_2025", "/api/v1/reports/tax/2025/summary"),
    ("tax_summary_before_2015", "/api/v1/reports/tax/2014/summary"),
    ("tax_summary_future", "/api/v1/reports/tax/2099/summary"),
    ("tax_pdf_2026", "/api/v1/reports/tax/2026/pdf"),
    ("tax_pdf_2025", "/api/v1/reports/tax/2025/pdf"),
    ("tax_pdf_future", "/api/v1/reports/tax/2099/pdf"),
    ("tax_excel_2026", "/api/v1/reports/tax/2026/excel"),
    ("tax_excel_2025", "/api/v1/reports/tax/2025/excel"),
    ("tax_excel_before_2015", "/api/v1/reports/tax/2014/excel"),
    ("securities_pdf", "/api/v1/reports/tax/securities/2025/pdf"),
    ("securities_excel", "/api/v1/reports/tax/securities/2025/excel"),
    ("transactions_pdf_all", "/api/v1/reports/transactions/pdf"),
    ("transactions_pdf_2025", "/api/v1/reports/transactions/pdf?year=2025"),
    ("transactions_excel_all", "/api/v1/reports/transactions/excel"),
    ("transactions_excel_2026", "/api/v1/reports/transactions/excel?year=2026"),
    ("transactions_csv_all", "/api/v1/reports/transactions/csv"),
    ("transactions_csv_2025", "/api/v1/reports/transactions/csv?year=2025"),
    ("available_years", "/api/v1/reports/available-years"),
]
EMPTY_ROUTES = [
    ("empty_performance_pdf", "GET", "/api/v1/reports/performance/pdf", None),
    ("empty_performance_excel", "GET", "/api/v1/reports/performance/excel", None),
    ("empty_tax_summary", "GET", "/api/v1/reports/tax/2026/summary", None),
    ("empty_tax_pdf", "GET", "/api/v1/reports/tax/2026/pdf", None),
    ("empty_tax_excel", "GET", "/api/v1/reports/tax/2026/excel", None),
    ("empty_transactions_pdf", "GET", "/api/v1/reports/transactions/pdf", None),
    ("empty_transactions_excel", "GET", "/api/v1/reports/transactions/excel", None),
    ("empty_transactions_csv", "GET", "/api/v1/reports/transactions/csv", None),
    ("empty_available_years", "GET", "/api/v1/reports/available-years", None),
    ("empty_rebalancing", "POST", "/api/v1/reports/rebalancing", TARGETS),
    ("empty_rebalancing_pdf", "POST", "/api/v1/reports/rebalancing/pdf", TARGETS),
    ("empty_drift", "GET", "/api/v1/reports/rebalancing/drift", None),
]
REBALANCING = [
    ("targets_before", "GET", "/api/v1/reports/rebalancing/targets", None),
    ("drift_before", "GET", "/api/v1/reports/rebalancing/drift", None),
    ("rebalancing", "POST", "/api/v1/reports/rebalancing", TARGETS),
    ("rebalancing_all_l1", "POST", "/api/v1/reports/rebalancing", {"L1": 1.0}),
    ("rebalancing_unknown_class", "POST", "/api/v1/reports/rebalancing", {"L1": 0.5, "NFT": 0.5}),
    ("rebalancing_bad_total", "POST", "/api/v1/reports/rebalancing", {"L1": 0.5, "Stable": 0.2}),
    ("rebalancing_pdf", "POST", "/api/v1/reports/rebalancing/pdf", TARGETS),
    ("targets_save_invalid", "PUT", "/api/v1/reports/rebalancing/targets", {"L1": 1.5}),
    ("targets_save", "PUT", "/api/v1/reports/rebalancing/targets", TARGETS),
    ("targets_after", "GET", "/api/v1/reports/rebalancing/targets", None),
    ("drift_after", "GET", "/api/v1/reports/rebalancing/drift", None),
]
OFFLINE = [
    ("offline_rebalancing", "POST", "/api/v1/reports/rebalancing", TARGETS),
    ("offline_performance_pdf", "GET", "/api/v1/reports/performance/pdf", None),
]


def _sorted_set(match: re.Match[str]) -> str:
    return "{" + ", ".join(sorted(match.group(1).split(", "))) + "}"


def _plain(value: Any) -> Any:
    if isinstance(value, dict):
        return {k: _plain(v) for k, v in sorted(value.items())}
    if isinstance(value, list | tuple):
        return [_plain(v) for v in value]
    if isinstance(value, float):
        return round(value, 4)
    if isinstance(value, datetime | date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, str):
        return SET_RE.sub(_sorted_set, UUID_RE.sub("<id>", value))
    return value


def _pdf_text(content: bytes) -> list[list[str]]:
    with fitz.open(stream=content, filetype="pdf") as document:
        return [[line for line in page.get_text().splitlines() if line.strip()] for page in document]


def _workbook(content: bytes) -> dict[str, Any]:
    workbook = load_workbook(io.BytesIO(content))
    return {
        sheet.title: [list(row) for row in sheet.iter_rows(values_only=True) if any(c is not None for c in row)]
        for sheet in workbook.worksheets
    }


def _body(response: Response) -> Any:
    kind = response.headers.get("content-type", "")
    if kind.startswith(PDF):
        return _pdf_text(response.content)
    if kind.startswith(XLSX):
        return _workbook(response.content)
    if kind.startswith("text/csv"):
        return list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";"))
    try:
        return response.json()
    except ValueError:
        return response.text


def _snapshot(response: Response) -> dict[str, Any]:
    headers = {k: response.headers[k] for k in ("content-type", "content-disposition") if k in response.headers}
    return _plain({"status": response.status_code, "headers": headers, "body": _body(response)})


def _price(symbol: str, day: int) -> Decimal:
    base, drift, amplitude, period, phase = SERIES[symbol]
    value = base * (1 + drift * day) * (1 + amplitude * math.sin(day / period + phase))
    return Decimal(str(round(value, 6)))


async def _quotes(symbols: Iterable[str], currency: str = "EUR") -> dict[str, PriceQuote]:
    return {s.upper(): PriceQuote(price=QUOTES[s.upper()]) for s in symbols if s.upper() in QUOTES}


async def _quote(symbol: str, asset_type: str, currency: str = "EUR") -> PriceQuote | None:
    price = QUOTES.get(symbol.upper())
    return PriceQuote(price=price) if price is not None else None


@pytest.fixture
def live_quotes(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(price_feed, "crypto", _quotes)
    monkeypatch.setattr(price_feed, "stocks", _quotes)
    monkeypatch.setattr(price_feed, "quote", _quote)


async def _seed_history(db: AsyncSession) -> None:
    start = FROZEN_NOW.date() - timedelta(days=HISTORY_DAYS)
    for symbol in SERIES:
        for day in range(HISTORY_DAYS + 1):
            moment: date = start + timedelta(days=day)
            db.add(AssetPriceHistory(symbol=symbol, price_date=moment, price_eur=_price(symbol, day), source="test"))
    await db.commit()


async def _seed_portfolio(client: AsyncClient, headers: dict[str, str]) -> None:
    portfolio = await client.post("/api/v1/portfolios", json={"name": "Rapports"}, headers=headers)
    assert portfolio.status_code == 201, portfolio.text
    portfolio_id = portfolio.json()["id"]
    asset_ids: dict[str, str] = {}
    for symbol, kind, exchange, price in ASSETS:
        body = {
            "portfolio_id": portfolio_id,
            "symbol": symbol,
            "asset_type": kind,
            "exchange": exchange,
            "current_price": price,
        }
        asset = await client.post("/api/v1/assets", json=body, headers=headers)
        assert asset.status_code == 201, asset.text
        asset_ids[symbol] = asset.json()["id"]
    for symbol, kind, quantity, price, fee, executed_at in TRANSACTIONS:
        body = {
            "asset_id": asset_ids[symbol],
            "transaction_type": kind,
            "quantity": quantity,
            "price": price,
            "fee": fee,
            "currency": "EUR",
            "executed_at": executed_at,
            "exchange": next(e for s, _, e, _ in ASSETS if s == symbol),
        }
        tx = await client.post("/api/v1/transactions", json=body, headers=headers)
        assert tx.status_code == 201, tx.text


def _compare(name: str, actual: dict[str, Any], mismatches: list[str]) -> None:
    path = GOLDEN_DIR / f"{name}.json"
    rendered = json.dumps(actual, indent=2, ensure_ascii=False, sort_keys=True) + "\n"
    if UPDATE or not path.exists():
        GOLDEN_DIR.mkdir(exist_ok=True)
        path.write_text(rendered, encoding="utf-8")
    elif path.read_text(encoding="utf-8") != rendered:
        mismatches.append(name)


async def _play(
    client: AsyncClient, scenarios: list[tuple[str, str, str, dict[str, str], dict[str, float] | None]]
) -> list[str]:
    mismatches: list[str] = []
    for name, method, url, who, body in scenarios:
        payload = {"target_allocations": body} if body is not None else None
        response = await client.request(method, url, headers=who, json=payload)
        _compare(name, _snapshot(response), mismatches)
    return mismatches


@pytest.mark.usefixtures("empty_test_redis", "offline", "live_quotes")
async def test_reports_routes_match_golden(
    client: AsyncClient, db_session: AsyncSession, regular_user: User, admin_user: User
) -> None:
    headers = {"Authorization": f"Bearer {create_token(str(regular_user.id))}"}
    empty = {"Authorization": f"Bearer {create_token(str(admin_user.id))}"}
    with time_machine.travel(FROZEN_NOW, tick=False):
        await _seed_history(db_session)
        await _seed_portfolio(client, headers)
        scenarios = [(name, "GET", url, headers, None) for name, url in GET_ROUTES]
        scenarios += [(name, method, url, headers, body) for name, method, url, body in REBALANCING]
        scenarios += [(name, method, url, empty, body) for name, method, url, body in EMPTY_ROUTES]
        scenarios.append(("unauthenticated", "GET", "/api/v1/reports/available-years", {}, None))
        mismatches = await _play(client, scenarios)
    assert not mismatches, f"golden reports différent : {mismatches} (GOLDEN_UPDATE=1 pour accepter)"


@pytest.mark.usefixtures("empty_test_redis", "offline")
async def test_reports_without_live_quotes_match_golden(
    client: AsyncClient, db_session: AsyncSession, regular_user: User
) -> None:
    headers = {"Authorization": f"Bearer {create_token(str(regular_user.id))}"}
    with time_machine.travel(FROZEN_NOW, tick=False):
        await _seed_history(db_session)
        await _seed_portfolio(client, headers)
        mismatches = await _play(client, [(name, method, url, headers, body) for name, method, url, body in OFFLINE])
    assert not mismatches, f"golden reports différent : {mismatches} (GOLDEN_UPDATE=1 pour accepter)"
