"""Filet de caractérisation des exports de transactions (PDF, Excel, CSV).

`report_transactions.py` était couvert à **18 %**. Ce sont trois fichiers que
l'utilisateur télécharge : un rapport faux s'y voit, mais aucun test ne le
regardait.

Les trois générateurs sont synchrones et ne prennent qu'un dictionnaire — ils
se testent en lisant le fichier produit, pas en simulant l'export.
"""

import csv
import io
from datetime import date
from types import SimpleNamespace

import fitz
import pytest
from openpyxl import load_workbook

from app.services.report_service import report_service


def tx(**surcharges):
    champs = {
        "date": date(2026, 3, 15),
        "transaction_type": "buy",
        "symbol": "BTC",
        "quantity": 0.5,
        "price": 40000.0,
        "total": 20000.0,
        "fee": 12.5,
    }
    champs.update(surcharges)
    return SimpleNamespace(**champs)


def texte_pdf(octets: bytes) -> str:
    with fitz.open(stream=octets, filetype="pdf") as doc:
        return "\n".join(page.get_text() for page in doc)


def lignes_csv(octets: bytes) -> list[list[str]]:
    texte = octets.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(texte), delimiter=";"))


def feuille_excel(octets: bytes):
    return load_workbook(io.BytesIO(octets)).active


class TestExportPdf:
    def test_le_pdf_porte_le_titre_et_la_periode(self):
        pdf = report_service.generate_transactions_pdf({"transactions": [tx()], "year": 2026})

        texte = texte_pdf(pdf)

        assert "Historique des Transactions" in texte
        assert "Année 2026" in texte

    def test_sans_annee_la_periode_couvre_tout(self):
        pdf = report_service.generate_transactions_pdf({"transactions": [tx()]})

        assert "Toutes les années" in texte_pdf(pdf)

    def test_le_resume_totalise_volume_et_frais(self):
        pdf = report_service.generate_transactions_pdf(
            {"transactions": [tx(total=20000.0, fee=12.5), tx(total=5000.0, fee=3.5)]}
        )

        texte = texte_pdf(pdf)

        assert "Nombre de transactions" in texte
        assert "25" in texte  # 25 000 EUR de volume, quel que soit le formatage
        assert "16" in texte  # 16 EUR de frais

    def test_un_export_vide_reste_un_pdf_lisible(self):
        """Zéro transaction ne produit ni erreur ni page blanche.

        Le résumé subsiste — « 0 transaction » — mais le tableau de détail est
        omis : l'utilisateur voit que l'export a fonctionné et qu'il n'y avait
        rien à exporter.
        """
        pdf = report_service.generate_transactions_pdf({"transactions": []})

        texte = texte_pdf(pdf)

        assert "Historique des Transactions" in texte
        assert "Détail" not in texte

    def test_les_types_sont_traduits_en_francais(self):
        pdf = report_service.generate_transactions_pdf({"transactions": [tx(transaction_type="staking_reward")]})

        assert "Reward" in texte_pdf(pdf)

    def test_un_type_inconnu_est_affiche_tel_quel(self):
        # `_TYPE_MAP.get(type, type)` : plutôt qu'une case vide, l'export montre
        # le libellé technique — visible, donc corrigeable.
        pdf = report_service.generate_transactions_pdf({"transactions": [tx(transaction_type="lending_reward")]})

        assert "lending_reward" in texte_pdf(pdf)


class TestExportExcel:
    def test_l_entete_couvre_les_sept_colonnes(self):
        ws = feuille_excel(report_service.generate_transactions_excel({"transactions": [tx()]}))

        entete = [c.value for c in ws[1]]

        assert entete == [
            "Date",
            "Type",
            "Actif",
            "Quantité",
            "Prix Unitaire (EUR)",
            "Valeur Totale (EUR)",
            "Frais (EUR)",
        ]

    def test_les_montants_sont_ecrits_comme_nombres_et_non_comme_texte(self):
        """Excel reçoit des flottants, pas des chaînes formatées.

        C'est ce qui permet à l'utilisateur de trier, sommer et filtrer. Le
        format d'affichage est porté par `number_format`, séparément de la
        valeur — contrairement à l'export CSV, qui aplatit tout en texte.
        """
        ws = feuille_excel(report_service.generate_transactions_excel({"transactions": [tx()]}))

        for colonne in (4, 5, 6, 7):
            valeur = ws.cell(row=2, column=colonne).value
            assert isinstance(valeur, (int, float)), f"colonne {colonne} écrite en texte"

        assert ws.cell(row=2, column=4).value == 0.5
        # 40000.0 se relit `40000` : openpyxl normalise les flottants entiers.
        assert ws.cell(row=2, column=5).value == 40000

    def test_la_quantite_garde_six_decimales_a_l_affichage(self):
        # Les fractions de crypto ne survivent pas à un format à deux décimales.
        ws = feuille_excel(report_service.generate_transactions_excel({"transactions": [tx()]}))

        assert ws.cell(row=2, column=4).number_format == "0.000000"

    def test_les_libelles_sont_plus_explicites_que_dans_le_pdf(self):
        # Le PDF, contraint par la largeur des colonnes, abrège en « Conv. ↓ ».
        excel = report_service.generate_transactions_excel({"transactions": [tx(transaction_type="conversion_in")]})

        assert feuille_excel(excel).cell(row=2, column=2).value == "Conversion entrante"

    def test_un_export_vide_ne_contient_que_l_entete(self):
        ws = feuille_excel(report_service.generate_transactions_excel({"transactions": []}))

        assert ws.max_row == 1


class TestExportCsv:
    def test_le_separateur_et_le_bom_visent_excel_francais(self):
        octets = report_service.generate_transactions_csv({"transactions": [tx()]})

        assert octets.startswith(b"\xef\xbb\xbf")  # BOM UTF-8
        assert b";" in octets

    def test_l_entete_est_identique_a_celui_de_l_excel(self):
        lignes = lignes_csv(report_service.generate_transactions_csv({"transactions": [tx()]}))

        assert lignes[0][0] == "Date"
        assert lignes[0][-1] == "Frais (EUR)"

    def test_les_nombres_utilisent_le_point_decimal(self):
        """Le CSV vise Excel français — BOM et point-virgule le disent — mais
        écrit ses nombres avec un **point** décimal.

        Excel en locale française attend la virgule : les colonnes Quantité,
        Prix, Valeur et Frais y arrivent en **texte**, non sommables et non
        triables. L'écart entre l'intention (le séparateur) et le format des
        nombres est épinglé ici, pas corrigé.
        """
        lignes = lignes_csv(report_service.generate_transactions_csv({"transactions": [tx()]}))

        assert lignes[1][3] == "0.500000"
        assert lignes[1][4] == "40000.00"
        assert "," not in lignes[1][4]

    def test_les_types_ne_sont_pas_traduits_contrairement_aux_deux_autres(self):
        """Troisième format, troisième traitement des types.

        Le PDF abrège (« Transfert ↓ »), l'Excel développe (« Transfert
        entrant »), le CSV n'traduit pas du tout et livre l'identifiant
        technique. Un utilisateur qui compare ses trois exports voit trois
        vocabulaires.
        """
        csv_octets = report_service.generate_transactions_csv({"transactions": [tx(transaction_type="transfer_in")]})
        excel = report_service.generate_transactions_excel({"transactions": [tx(transaction_type="transfer_in")]})

        assert lignes_csv(csv_octets)[1][1] == "transfer_in"
        assert feuille_excel(excel).cell(row=2, column=2).value == "Transfert entrant"

    def test_un_export_vide_ne_contient_que_l_entete(self):
        lignes = lignes_csv(report_service.generate_transactions_csv({"transactions": []}))

        assert len(lignes) == 1


class TestDateManquante:
    @pytest.mark.parametrize("generateur", ["pdf", "excel", "csv"])
    def test_une_transaction_sans_date_laisse_la_cellule_vide(self, generateur):
        # `tx.date.strftime(...) if tx.date else ""` : les trois formats
        # préfèrent une case vide à une erreur d'export.
        donnees = {"transactions": [tx(date=None)]}

        if generateur == "pdf":
            assert "BTC" in texte_pdf(report_service.generate_transactions_pdf(donnees))
        elif generateur == "excel":
            # openpyxl relit une chaîne vide comme `None`, pas comme "".
            assert (
                feuille_excel(report_service.generate_transactions_excel(donnees)).cell(row=2, column=1).value is None
            )
        else:
            assert lignes_csv(report_service.generate_transactions_csv(donnees))[1][0] == ""
