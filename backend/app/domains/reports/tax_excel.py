import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.domains.reports.styles import XL_BORDER, XL_HEADER_FILL, XL_HEADER_FONT, XL_MONEY, now_paris
from app.domains.reports.tax import TaxEvent2086, TaxSummary2086

_EVENT_HEADERS = [
    "Date",
    "Actif",
    "Type",
    "Quantité",
    "Prix unitaire",
    "Prix cession",
    "Valeur portefeuille",
    "Coût acquisition",
    "Fraction acq.",
    "PV/MV",
    "Durée",
]
_MONEY_COLS = (5, 6, 7, 8, 9, 10)
_GAIN_COL = 10


def _summary_items(tax: TaxSummary2086) -> list[tuple[str, Any, str | None]]:
    return [
        ("Nombre de cessions", tax.nb_cessions, None),
        ("  dont court terme (< 2 ans)", tax.nb_court_terme, None),
        ("  dont long terme (≥ 2 ans)", tax.nb_long_terme, None),
        ("Prix total de cession", tax.total_cessions, XL_MONEY),
        ("Fraction d'acquisition (PMP)", tax.total_acquisitions_fraction, XL_MONEY),
        ("Total plus-values", tax.total_plus_values, XL_MONEY),
        ("Total moins-values", tax.total_moins_values, XL_MONEY),
        ("Plus-value nette imposable", tax.net_plus_value, XL_MONEY),
    ]


def _summary_sheet(ws: Worksheet, tax: TaxSummary2086, year: int) -> None:
    ws.title = f"Résumé {year}"
    ws["A1"] = f"Déclaration Fiscale Crypto — Année {year}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Généré le {now_paris().strftime('%d/%m/%Y')}"
    ws["A2"].font = Font(italic=True, color="666666")
    ws["A4"] = "Résumé des cessions"
    ws["A4"].font = Font(bold=True, size=12)
    summary_items = _summary_items(tax)
    for i, (label, value, fmt) in enumerate(summary_items, 5):
        ws[f"A{i}"] = label
        ws[f"B{i}"] = value
        if fmt:
            ws[f"B{i}"].number_format = fmt
    row = 5 + len(summary_items) + 1
    ws[f"A{row}"] = "Estimation fiscale (PFU 30%)"
    ws[f"A{row}"].font = Font(bold=True, size=12)
    tax_items = [
        ("Impôt sur le revenu (12,8%)", tax.ir_12_8),
        ("Prélèvements sociaux (17,2%)", tax.ps_17_2),
        ("Total Flat Tax (30%)", tax.flat_tax_30),
    ]
    for i, (label, value) in enumerate(tax_items, row + 1):
        ws[f"A{i}"] = label
        ws[f"B{i}"] = value
        ws[f"B{i}"].number_format = XL_MONEY
        if label.startswith("Total"):
            ws[f"A{i}"].font = Font(bold=True)
            ws[f"B{i}"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 18


def _event_values(e: TaxEvent2086) -> list[Any]:
    return [
        e.date.strftime("%d/%m/%Y"),
        e.symbol,
        "Conversion" if e.event_type == "conversion_out" else "Vente",
        e.quantity,
        e.unit_price,
        e.cession_price,
        e.portfolio_value,
        e.total_acquisition_cost,
        e.acquisition_fraction,
        e.gain_loss,
        "≥ 2 ans" if e.holding_period == "long_terme" else "< 2 ans",
    ]


def _events_sheet(wb: Workbook, events: list[TaxEvent2086]) -> None:
    ws = wb.create_sheet("Cessions")
    for col, h in enumerate(_EVENT_HEADERS, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = XL_HEADER_FONT
        c.fill = XL_HEADER_FILL
        c.border = XL_BORDER
    for row, e in enumerate(sorted(events, key=lambda x: x.date), 2):
        for col, v in enumerate(_event_values(e), 1):
            c = ws.cell(row=row, column=col, value=v)
            c.border = XL_BORDER
            if col in _MONEY_COLS:
                c.number_format = XL_MONEY
            if col == _GAIN_COL:
                c.font = Font(color="16a34a" if e.gain_loss >= 0 else "dc2626")
    for col in range(1, 12):
        ws.column_dimensions[get_column_letter(col)].width = 16


def render_tax_excel(tax: TaxSummary2086, year: int) -> bytes:
    wb = Workbook()
    _summary_sheet(wb.active, tax, year)
    _events_sheet(wb, tax.events)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
