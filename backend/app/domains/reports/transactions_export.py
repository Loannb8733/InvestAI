import csv
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Flowable, Paragraph, Spacer, Table, TableStyle

from app.domains.reports.data import TaxTransaction
from app.domains.reports.styles import (
    BLUE,
    BORDER_COLOR,
    DARK_BLUE,
    LIGHT_BG,
    XL_BORDER,
    XL_HEADER_FILL,
    XL_HEADER_FONT,
    XL_MONEY,
    build_doc,
    fmt_qty,
    format_paris_dt,
    money,
)

_PDF_TYPE_LABELS = {
    "buy": "Achat",
    "sell": "Vente",
    "transfer_in": "Transfert ↓",
    "transfer_out": "Transfert ↑",
    "conversion_in": "Conv. ↓",
    "conversion_out": "Conv. ↑",
    "airdrop": "Airdrop",
    "staking_reward": "Reward",
    "fee": "Frais",
    "dividend": "Dividende",
    "interest": "Intérêt",
    "staking": "Staking",
    "unstaking": "Unstaking",
}
_EXCEL_TYPE_LABELS = {
    "buy": "Achat",
    "sell": "Vente",
    "transfer_in": "Transfert entrant",
    "transfer_out": "Transfert sortant",
    "conversion_in": "Conversion entrante",
    "conversion_out": "Conversion sortante",
    "airdrop": "Airdrop",
    "staking_reward": "Reward",
    "fee": "Frais",
    "dividend": "Dividende",
    "interest": "Intérêt",
    "staking": "Staking",
    "unstaking": "Unstaking",
}
_EXPORT_HEADERS = [
    "Date",
    "Type",
    "Actif",
    "Quantité",
    "Prix Unitaire (EUR)",
    "Valeur Totale (EUR)",
    "Frais (EUR)",
]
_PDF_DETAIL_COL_WIDTHS = [2.2, 2.2, 2, 2.4, 2.8, 2.8, 2]


def _date_label(tx: TaxTransaction) -> str:
    return tx.date.strftime("%d/%m/%Y") if tx.date else ""


def _pdf_summary_section(transactions: list[TaxTransaction], heading_style: ParagraphStyle) -> list[Flowable]:
    total_volume = sum(t.total for t in transactions)
    total_fees = sum(t.fee for t in transactions)
    rows = [
        ["Nombre de transactions", str(len(transactions))],
        ["Volume total", money(total_volume)],
        ["Total frais", money(total_fees)],
    ]
    t = Table(rows, colWidths=[8 * cm, 5 * cm])
    t.setStyle(
        TableStyle(
            [
                ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                ("FONTSIZE", (0, 0), (-1, -1), 10),
                ("GRID", (0, 0), (-1, -1), 1, BORDER_COLOR),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("BACKGROUND", (0, 0), (-1, -1), LIGHT_BG),
            ]
        )
    )
    return [Paragraph("Résumé", heading_style), t, Spacer(1, 15)]


def _pdf_detail_section(transactions: list[TaxTransaction], heading_style: ParagraphStyle) -> list[Flowable]:
    if not transactions:
        return []
    rows = [["Date", "Type", "Actif", "Quantité", "Prix Unitaire", "Valeur Totale", "Frais"]]
    for tx in transactions:
        rows.append(
            [
                _date_label(tx),
                _PDF_TYPE_LABELS.get(tx.transaction_type, tx.transaction_type),
                tx.symbol,
                fmt_qty(tx.quantity),
                money(tx.price),
                money(tx.total),
                money(tx.fee),
            ]
        )
    t = Table(rows, colWidths=[w * cm for w in _PDF_DETAIL_COL_WIDTHS])
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), BLUE),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (3, 0), (-1, -1), "RIGHT"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 1, BORDER_COLOR),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT_BG]),
            ]
        )
    )
    return [Paragraph("Détail", heading_style), t]


def render_transactions_pdf(data: dict[str, Any]) -> bytes:
    transactions = data.get("transactions", [])
    buffer = io.BytesIO()
    doc = build_doc(
        buffer,
        rightMargin=1.5 * cm,
        leftMargin=1.5 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("T", parent=styles["Heading1"], fontSize=20, spaceAfter=20, textColor=BLUE)
    heading_style = ParagraphStyle(
        "H",
        parent=styles["Heading2"],
        fontSize=12,
        spaceBefore=15,
        spaceAfter=8,
        textColor=DARK_BLUE,
    )
    year = data.get("year")
    period = f"Année {year}" if year else "Toutes les années"
    elements: list[Flowable] = [
        Paragraph("Historique des Transactions", title_style),
        Paragraph(f"{period} — Généré le {format_paris_dt()}", styles["Normal"]),
        Spacer(1, 15),
    ]
    elements += _pdf_summary_section(transactions, heading_style)
    elements += _pdf_detail_section(transactions, heading_style)
    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()


def render_transactions_excel(data: dict[str, Any]) -> bytes:
    transactions: list[TaxTransaction] = data.get("transactions", [])
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    for col, h in enumerate(_EXPORT_HEADERS, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = XL_HEADER_FONT
        c.fill = XL_HEADER_FILL
        c.border = XL_BORDER
    for row, tx in enumerate(transactions, 2):
        ws.cell(row=row, column=1, value=_date_label(tx)).border = XL_BORDER
        ws.cell(
            row=row,
            column=2,
            value=_EXCEL_TYPE_LABELS.get(tx.transaction_type, tx.transaction_type),
        ).border = XL_BORDER
        ws.cell(row=row, column=3, value=tx.symbol).border = XL_BORDER
        c = ws.cell(row=row, column=4, value=tx.quantity)
        c.number_format = "0.000000"
        c.border = XL_BORDER
        for col, value in ((5, tx.price), (6, tx.total), (7, tx.fee)):
            ws.cell(row=row, column=col, value=value).number_format = XL_MONEY
            ws.cell(row=row, column=col).border = XL_BORDER
    for col in range(1, len(_EXPORT_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _french_number(value: float, decimals: int) -> str:
    return f"{value:.{decimals}f}".replace(".", ",")


def render_transactions_csv(data: dict[str, Any]) -> bytes:
    """CSV pour Excel en locale française : BOM utf-8-sig, point-virgule, virgule décimale sans séparateur de milliers.

    Avec un point décimal, Excel français lit les colonnes chiffrées comme du texte :
    elles ne se somment ni ne se trient.
    """
    transactions: list[TaxTransaction] = data.get("transactions", [])
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(_EXPORT_HEADERS)
    for tx in transactions:
        writer.writerow(
            [
                _date_label(tx),
                tx.transaction_type,
                tx.symbol,
                _french_number(tx.quantity, 6),
                _french_number(tx.price, 2),
                _french_number(tx.total, 2),
                _french_number(tx.fee, 2),
            ]
        )
    return buffer.getvalue().encode("utf-8-sig")
