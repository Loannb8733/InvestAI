import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.domains.reports.data import AssetReport, PortfolioSummary, TaxTransaction
from app.domains.reports.styles import (
    XL_BORDER,
    XL_HEADER_FILL,
    XL_HEADER_FONT,
    XL_MONEY,
    XL_PERCENT,
    format_paris_dt,
)

_ASSET_HEADERS = [
    *("Symbole", "Nom", "Type", "Plateforme", "Quantité", "PRU", "Prix actuel", "Break-even"),
    *("Investi", "Valeur", "+/- Value", "Perf.", "Frais", "Risque"),
]
_ASSET_MONEY_COLS = {6, 7, 8, 9, 10, 11, 13}
_ASSET_PERCENT_COLS = {12, 14}
_PLATFORM_HEADERS = ["Plateforme", "Nb Actifs", "Valeur", "Investi", "Frais", "P&L Net", "ROI %"]

LabelledValue = tuple[str, Any, str | None]


def _write_headers(ws: Worksheet, row: int, headers: list[str]) -> None:
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = XL_HEADER_FONT
        c.fill = XL_HEADER_FILL
        c.border = XL_BORDER


def _summary_rows(summary: dict[str, Any], pnl: dict[str, Any]) -> list[LabelledValue]:
    rows: list[LabelledValue] = [
        ("Valeur totale", summary.get("total_value", 0), XL_MONEY),
        ("Total investi", summary.get("total_invested", 0), XL_MONEY),
        ("Plus/Moins-value", summary.get("gain_loss", 0), XL_MONEY),
        ("Performance (%)", summary.get("gain_loss_percent", 0) / 100, XL_PERCENT),
    ]
    if pnl:
        rows.extend(
            [
                ("", "", None),
                ("P&L Latent", pnl.get("unrealized_pnl", 0), XL_MONEY),
                ("P&L Réalisé", pnl.get("realized_pnl", 0), XL_MONEY),
                ("Total Frais", pnl.get("total_fees", 0), XL_MONEY),
                ("P&L Net", pnl.get("net_pnl", 0), XL_MONEY),
            ]
        )
    return rows


def _summary_sheet(ws: Worksheet, summary: dict[str, Any], pnl: dict[str, Any]) -> None:
    ws.title = "Résumé"
    ws["A1"] = "Rapport de Performance InvestAI"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"Généré le {format_paris_dt()}"
    for cell_ref in ["A4", "B4"]:
        ws[cell_ref].font = XL_HEADER_FONT
        ws[cell_ref].fill = XL_HEADER_FILL
    ws["A4"] = "Indicateur"
    ws["B4"] = "Valeur"
    rows = _summary_rows(summary, pnl)
    for i, (label, value, fmt) in enumerate(rows, start=5):
        ws[f"A{i}"] = label
        if value != "":
            ws[f"B{i}"] = value
        if fmt:
            ws[f"B{i}"].number_format = fmt
    if pnl:
        net_row = 5 + len(rows) - 1
        ws[f"A{net_row}"].font = Font(bold=True)
        ws[f"B{net_row}"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 18


def _portfolios_sheet(wb: Workbook, portfolios: list[PortfolioSummary]) -> None:
    if not portfolios:
        return
    ws = wb.create_sheet("Portefeuilles")
    _write_headers(ws, 1, ["Nom", "Valeur", "Investi", "+/- Value", "Performance", "Nb Actifs"])
    for row, p in enumerate(portfolios, 2):
        ws.cell(row=row, column=1, value=p.name).border = XL_BORDER
        ws.cell(row=row, column=2, value=p.total_value).number_format = XL_MONEY
        ws.cell(row=row, column=3, value=p.total_invested).number_format = XL_MONEY
        ws.cell(row=row, column=4, value=p.gain_loss).number_format = XL_MONEY
        ws.cell(row=row, column=5, value=p.gain_loss_percent / 100).number_format = XL_PERCENT
        ws.cell(row=row, column=6, value=p.asset_count)
        for col in range(1, 7):
            ws.cell(row=row, column=col).border = XL_BORDER
    for col in range(1, 7):
        ws.column_dimensions[get_column_letter(col)].width = 15


def _asset_values(a: AssetReport) -> list[Any]:
    return [
        a.symbol,
        a.name,
        a.asset_type,
        a.exchange or "",
        a.quantity,
        a.avg_buy_price,
        a.current_price,
        a.breakeven_price if a.breakeven_price else "",
        a.total_invested,
        a.current_value,
        a.gain_loss,
        a.gain_loss_percent / 100,
        a.total_fees,
        a.risk_weight / 100 if a.risk_weight else "",
    ]


def _assets_sheet(wb: Workbook, assets: list[AssetReport]) -> None:
    if not assets:
        return
    ws = wb.create_sheet("Actifs")
    _write_headers(ws, 1, _ASSET_HEADERS)
    for row, a in enumerate(assets, 2):
        for col, v in enumerate(_asset_values(a), 1):
            c = ws.cell(row=row, column=col, value=v)
            c.border = XL_BORDER
            if col in _ASSET_MONEY_COLS and v != "":
                c.number_format = XL_MONEY
            elif col in _ASSET_PERCENT_COLS and v != "":
                c.number_format = XL_PERCENT
    for col in range(1, len(_ASSET_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14


def _transactions_sheet(wb: Workbook, transactions: list[TaxTransaction]) -> None:
    if not transactions:
        return
    ws = wb.create_sheet("Transactions")
    _write_headers(ws, 1, ["Date", "Symbole", "Type", "Quantité", "Prix", "Total", "Frais"])
    for row, t in enumerate(transactions, 2):
        date_str = t.date.strftime("%d/%m/%Y") if t.date else ""
        ws.cell(row=row, column=1, value=date_str).border = XL_BORDER
        ws.cell(row=row, column=2, value=t.symbol).border = XL_BORDER
        ws.cell(row=row, column=3, value=t.transaction_type).border = XL_BORDER
        ws.cell(row=row, column=4, value=t.quantity).border = XL_BORDER
        ws.cell(row=row, column=5, value=t.price).number_format = XL_MONEY
        ws.cell(row=row, column=6, value=t.total).number_format = XL_MONEY
        ws.cell(row=row, column=7, value=t.fee).number_format = XL_MONEY
        for col in range(1, 8):
            ws.cell(row=row, column=col).border = XL_BORDER
    for col in range(1, 8):
        ws.column_dimensions[get_column_letter(col)].width = 14


def _risk_rows(risk: dict[str, Any]) -> list[LabelledValue]:
    mdd = risk.get("max_drawdown", {})
    var95 = risk.get("var_95", {})
    conc = risk.get("concentration", {})
    st20 = risk.get("stress_test_20", {})
    st40 = risk.get("stress_test_40", {})
    return [
        ("Volatilité annualisée", risk.get("volatility", 0) / 100, XL_PERCENT),
        ("Ratio de Sharpe", risk.get("sharpe_ratio", 0), "0.00"),
        ("Max Drawdown", mdd.get("max_drawdown_percent", 0) / 100, XL_PERCENT),
        ("Max Drawdown — Pic", mdd.get("peak_date", ""), None),
        ("Max Drawdown — Creux", mdd.get("trough_date", ""), None),
        ("VaR 95% (%)", var95.get("var_percent", 0) / 100, XL_PERCENT),
        ("VaR 95% (montant)", var95.get("var_amount", 0), XL_MONEY),
        ("", "", None),
        ("Indice HHI", conc.get("hhi", 0), "0"),
        ("Concentration", conc.get("interpretation", ""), None),
        ("Actif dominant", conc.get("top_asset", ""), None),
        ("Poids actif dominant", (conc.get("top_concentration", 0) or 0) / 100, XL_PERCENT),
        ("", "", None),
        ("Stress Test -20% (perte)", st20.get("potential_loss", 0), XL_MONEY),
        ("Stress Test -20% (valeur)", st20.get("stressed_value", 0), XL_MONEY),
        ("Stress Test -40% (perte)", st40.get("potential_loss", 0), XL_MONEY),
        ("Stress Test -40% (valeur)", st40.get("stressed_value", 0), XL_MONEY),
    ]


def _risk_sheet(wb: Workbook, risk: dict[str, Any]) -> None:
    if not risk:
        return
    ws = wb.create_sheet("Analyse de Risque")
    _write_headers(ws, 1, ["Indicateur", "Valeur"])
    for i, (label, value, fmt) in enumerate(_risk_rows(risk), start=2):
        ws[f"A{i}"] = label
        ws[f"A{i}"].border = XL_BORDER
        if value != "":
            ws[f"B{i}"] = value
        ws[f"B{i}"].border = XL_BORDER
        if fmt:
            ws[f"B{i}"].number_format = fmt
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 18


def _attribution_sheet(wb: Workbook, attr: dict[str, float]) -> None:
    if not any(v > 0 for v in attr.values()):
        return
    ws = wb.create_sheet("Attribution")
    ws["A1"] = "Attribution de Performance"
    ws["A1"].font = Font(bold=True, size=14)
    _write_headers(ws, 3, ["Catégorie", "Valeur", "Poids"])
    total_attr = sum(attr.values()) or 1
    rows = [
        ("Alpha (Altcoins)", attr.get("alpha", 0)),
        ("Beta (BTC)", attr.get("beta", 0)),
        ("Protection (Or)", attr.get("protection", 0)),
        ("Munitions (Cash/Stables)", attr.get("munitions", 0)),
    ]
    for i, (label, value) in enumerate(rows, 4):
        ws.cell(row=i, column=1, value=label).border = XL_BORDER
        c_val = ws.cell(row=i, column=2, value=value)
        c_val.number_format = XL_MONEY
        c_val.border = XL_BORDER
        c_pct = ws.cell(row=i, column=3, value=value / total_attr)
        c_pct.number_format = XL_PERCENT
        c_pct.border = XL_BORDER
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 12


def _cash_flow_sheet(wb: Workbook, cf: dict[str, Any], summary: dict[str, Any]) -> None:
    if not (cf.get("total_deposits", 0) > 0 or cf.get("total_withdrawals", 0) > 0):
        return
    ws = wb.create_sheet("Flux de Trésorerie")
    ws["A1"] = "Flux de Trésorerie (Cash Flows)"
    ws["A1"].font = Font(bold=True, size=14)
    _write_headers(ws, 3, ["Flux", "Montant"])
    rows = [
        ("Total déposé", cf.get("total_deposits", 0)),
        ("Total retiré", cf.get("total_withdrawals", 0)),
        ("Flux nets (dépôts − retraits)", cf.get("net_flows", 0)),  # noqa: RUF001
        ("Plus/Moins-value réelle", summary.get("gain_loss", 0)),
    ]
    for i, (label, value) in enumerate(rows, 4):
        ws.cell(row=i, column=1, value=label).border = XL_BORDER
        c = ws.cell(row=i, column=2, value=value)
        c.number_format = XL_MONEY
        c.border = XL_BORDER
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 18


def _platforms_sheet(wb: Workbook, platforms: list[dict[str, Any]]) -> None:
    if not platforms:
        return
    ws = wb.create_sheet("Perf. par Plateforme")
    _write_headers(ws, 1, _PLATFORM_HEADERS)
    for row, p in enumerate(platforms, 2):
        vals = [
            p["exchange"],
            p["count"],
            p["total_value"],
            p["total_invested"],
            p["total_fees"],
            p["net_pnl"],
            p["roi_percent"] / 100,
        ]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.border = XL_BORDER
            if col in (3, 4, 5, 6):
                c.number_format = XL_MONEY
            elif col == 7:
                c.number_format = XL_PERCENT
    for col in range(1, len(_PLATFORM_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16


def render_performance_excel(data: dict[str, Any]) -> bytes:
    wb = Workbook()
    summary = data.get("summary", {})
    _summary_sheet(wb.active, summary, data.get("pnl_data", {}))
    _portfolios_sheet(wb, data.get("portfolios", []))
    _assets_sheet(wb, data.get("assets", []))
    _transactions_sheet(wb, data.get("transactions", []))
    _risk_sheet(wb, data.get("risk_metrics", {}))
    _attribution_sheet(wb, data.get("attribution", {}))
    _cash_flow_sheet(wb, data.get("cash_flows", {}), summary)
    _platforms_sheet(wb, data.get("platform_analysis", []))
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
